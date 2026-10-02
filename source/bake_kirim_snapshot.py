"""
Kirim/Chiqim jurnali uchun zaxira nusxa: ASTATKA jadvalidagi "AppSheet_Kirim" varag'ini
quyish-open.html ichidagi <template id="kirim-snap"> ga CSV ko'rinishida joylaydi.

Nima uchun: claude.ai artifact sahifasi ba'zi brauzerlarda Google Sheets'ga to'g'ridan-to'g'ri
ulana olmaydi. Shunda sahifa jonli ma'lumot o'rniga shu nusxani ko'rsatadi ("... holatidagi nusxa").

Ishlatilishi:
    python3 bake_kirim_snapshot.py quyish-open.html /tmp/astatka.xlsx

Faqat <template id="kirim-snap" ...>...</template> blokini almashtiradi; boshqa hech narsaga tegmaydi.
"""
import sys, re, csv, io, datetime, html
import openpyxl

def cell(v):
    if v is None: return ''
    if isinstance(v, datetime.datetime): return v.strftime('%d/%m/%Y %H:%M:%S')
    if isinstance(v, datetime.date): return v.strftime('%d/%m/%Y')
    if isinstance(v, float) and v == int(v): return str(int(v))
    return str(v)

def main():
    html_path, xlsx_path = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if 'AppSheet_Kirim' not in wb.sheetnames:
        raise SystemExit('AppSheet_Kirim varag\'i topilmadi — hech narsa o\'zgartirilmadi')
    ws = wb['AppSheet_Kirim']
    buf = io.StringIO(); w = csv.writer(buf, quoting=csv.QUOTE_ALL, lineterminator='\n')
    n = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        vals = [cell(v) for v in row]
        if i > 0 and not any(vals): continue
        w.writerow(vals); n += i > 0
    now = datetime.datetime.utcnow() + datetime.timedelta(hours=5)  # Toshkent
    src = open(html_path, encoding='utf-8').read()
    pat = re.compile(r'<template id="kirim-snap"[^>]*>.*?</template>', re.S)
    if len(pat.findall(src)) != 1:
        raise SystemExit('kirim-snap bloki topilmadi yoki bir nechta — hech narsa o\'zgartirilmadi')
    block = '<template id="kirim-snap" data-at="%s">%s</template>' % (now.strftime('%d.%m.%Y %H:%M'), html.escape(buf.getvalue(), quote=False))
    src = pat.sub(lambda m: block, src)
    open(html_path, 'w', encoding='utf-8').write(src)
    print('OK: AppSheet_Kirim nusxasi joylandi (%d ta yozuv, %s)' % (n, now.strftime('%d.%m.%Y %H:%M')))

if __name__ == '__main__':
    main()
