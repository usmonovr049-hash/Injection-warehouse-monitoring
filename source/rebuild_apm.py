"""
APM bo'limini "Zaxira Yetish Tahlili (IMM Grafik)" jadvalidan yangilash (09.10.2026'dan).

Ishlatilishi:
    python3 rebuild_apm.py apm/apm.js /tmp/zyt.xlsx

Manba jadval: fileId 1I4GbqRL-r9_0pDiIL0hQatOSv_dAriYT5XbMr5K8Vv4 (xlsx qilib yuklab oling).
O'qiladigan varaqlar:
  - "Fakt"          : 1-qator sarlavha (A raqam, B nom, C.. kunlar sanasi), har kun ustunida detal donalari
  - "Fact smena"    : sarlavha "01.10 1sm", "01.10 2sm", ... (smenalar)
  - "Berilgan_Plan" : "Fakt" kabi, kunlik reja (bo'sh katak = reja kiritilmagan)
  - "IMM_Yuklama"   : B2 sana, 6-15 qatorlar: IMM, tonnaj, plan soat, mavjud soat, ... , fakt soat (F)
  - "IMM_Taqsimot"  : B2 sana, 6-15 qatorlar: IMM, Detal-1, Dona-1, Detal-2, Dona-2, Detal-3, Dona-3

apm.js ichidagi "// APM_DATA_START" ... "// APM_DATA_END" orasidagi APM_MONTHS massivini yangilaydi:
  - joriy oy oxirgi element bo'lsa — uni yangilaydi (IMM kunlik nusxalari saqlanib, yangisi qo'shiladi);
  - yangi oy boshlansa — oxirgisini frozen:true qiladi va yangi oy qo'shadi;
  - muzlatilgan oyni yoki kunlar soni kamayishini XAVFLI deb to'xtaydi (hech narsa o'zgarmaydi).
"""
import sys, re, json, datetime
import openpyxl

UZ = {1:('Yanvar','yan'),2:('Fevral','fev'),3:('Mart','mar'),4:('Aprel','apr'),5:('May','may'),6:('Iyun','iyun'),
      7:('Iyul','iyul'),8:('Avgust','avg'),9:('Sentabr','sen'),10:('Oktabr','okt'),11:('Noyabr','noy'),12:('Dekabr','dek')}

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None

def daily(ws):
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    cols = [(i, h) for i, h in enumerate(head) if i >= 2 and isinstance(h, datetime.datetime)]
    data = [r for r in rows[1:] if r and r[0] is not None]
    out = {}
    for i, h in cols:
        vals = [num(r[i]) for r in data if i < len(r)]
        vals = [v for v in vals if v is not None]
        out[h.date()] = (sum(vals) if vals else None)
    return out

def main():
    js_path, xlsx = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    fakt = daily(wb['Fakt']); plan = daily(wb['Berilgan_Plan'])
    if not fakt: raise SystemExit('"Fakt" varag\'ida sana ustunlari topilmadi — hech narsa o\'zgartirilmadi')
    first = min(fakt); y, m = first.year, first.month
    mid = '%04d-%02d' % (y, m)
    # smenalar
    sm = {}
    rows = list(wb['Fact smena'].iter_rows(values_only=True))
    head = rows[0]; data = [r for r in rows[1:] if r and r[0] is not None]
    for i, h in enumerate(head):
        mm = re.match(r'^\s*(\d{1,2})\.(\d{1,2})\s*([12])\s*sm', str(h or ''))
        if not mm or int(mm.group(2)) != m: continue
        d, s = int(mm.group(1)), int(mm.group(3))
        sm[(d, s)] = sum(v for v in (num(r[i]) for r in data if i < len(r)) if v)
    month_days = [d for d in fakt if d.year == y and d.month == m]
    def nz(d):
        return (fakt.get(d) or 0) > 0 or (plan.get(d) or 0) > 0 or sm.get((d.day, 1), 0) > 0 or sm.get((d.day, 2), 0) > 0
    active = [d.day for d in month_days if nz(d)]
    if not active: raise SystemExit('Oy uchun hech qanday fakt/reja yo\'q — hech narsa o\'zgartirilmadi')
    N = max(active)
    def day(d): return datetime.date(y, m, d)
    entry = {
        'id': mid, 'label': UZ[m][0], 'short': UZ[m][1],
        'asOf': (datetime.datetime.utcnow() + datetime.timedelta(hours=5)).strftime('%d.%m.%Y'),
        'frozen': False, 'days': N,
        's1': [round(sm.get((d, 1), 0)) for d in range(1, N + 1)],
        's2': [round(sm.get((d, 2), 0)) for d in range(1, N + 1)],
        'fakt': [round(fakt.get(day(d)) or 0) for d in range(1, N + 1)],
        'plan': [(None if plan.get(day(d)) is None else round(plan[day(d)])) for d in range(1, N + 1)],
        'imm': {},
    }
    # IMM kunlik nusxasi
    yk, tq = wb['IMM_Yuklama'], wb['IMM_Taqsimot']
    ydate = list(yk.iter_rows(min_row=2, max_row=2, values_only=True))[0][1]
    snap = {}
    if isinstance(ydate, datetime.datetime) and ydate.year == y and ydate.month == m:
        names = {}
        for r in wb['Maxsulotlar'].iter_rows(min_row=2, values_only=True):
            if r and r[1] is not None: names[str(r[1]).split('.')[0]] = str(r[2] or '')
        parts = {}
        for r in tq.iter_rows(min_row=6, max_row=15, values_only=True):
            if not r or not r[1]: continue
            p = []
            for k in (2, 4, 6):
                if r[k] is not None and (num(r[k+1]) or 0) > 0:
                    code = str(r[k]).split('.')[0]
                    p.append('%s %s — %d' % (code, names.get(code, '')[:24], round(num(r[k+1]))))
            parts[str(r[1]).strip()] = '; '.join(p)
        for r in yk.iter_rows(min_row=6, max_row=15, values_only=True):
            if not r or not r[0] or not str(r[0]).startswith('№'): continue
            snap[str(r[0]).strip()] = {'plan': round(num(r[2]) or 0, 1), 'fact': round(num(r[5]) or 0, 1), 'parts': parts.get(str(r[0]).strip(), '')}
        if sum(v['plan'] for v in snap.values()) > 0:
            entry['imm'][str(ydate.day)] = snap
    # apm.js ga yozish
    src = open(js_path, encoding='utf-8').read()
    a = src.index('// APM_DATA_START'); b = src.index('// APM_DATA_END')
    block = src[a:b]
    arr = json.loads(block[block.index('['):block.rindex(']') + 1])
    ids = [x['id'] for x in arr]
    if mid in ids[:-1] or (ids and mid < ids[-1]):
        raise SystemExit('XAVFLI: %s oyi allaqachon arxivda yoki oxirgi oydan oldin — hech narsa o\'zgartirilmadi' % mid)
    if ids and ids[-1] == mid:
        old = arr[-1]
        if N < old['days']:
            raise SystemExit('XAVFLI: %s uchun kunlar soni %d dan %d ga kamaydi — hech narsa o\'zgartirilmadi' % (mid, old['days'], N))
        merged = dict(old.get('imm') or {}); merged.update(entry['imm']); entry['imm'] = merged
        arr[-1] = entry; msg = 'Mavjud oy yangilandi'
    else:
        if arr: arr[-1]['frozen'] = True
        arr.append(entry); msg = 'YANGI OY qo\'shildi (oldingi oy muzlatildi)'
    body = 'var APM_MONTHS=[\n' + ',\n'.join(json.dumps(x, ensure_ascii=False, separators=(',', ':')) for x in arr) + '\n];\n'
    src = src[:a] + '// APM_DATA_START\n' + body + src[b:]
    open(js_path, 'w', encoding='utf-8').write(src)
    ptot = sum(p for p in entry['plan'] if p)
    print('OK: %s — %s: %d kun, fakt %d, reja %d (%d kun), IMM kunlari: %s' % (
        mid, msg, N, sum(entry['fakt']), ptot, sum(1 for p in entry['plan'] if p), ','.join(sorted(entry['imm'], key=int)) or '-'))

if __name__ == '__main__':
    main()
